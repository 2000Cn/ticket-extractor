#!/usr/bin/env python3
"""
Flight Ticket Extractor - friendly desktop app (100% local, no API).

For END USERS:  just double-click the packaged app, click "Add ticket PDFs",
then "Create Excel". No Python knowledge needed.

For YOU (to build the app):  see the README / packaging steps provided
separately. In short:  pip install pyinstaller pdfplumber openpyxl
then:  pyinstaller --onefile --windowed --name "TicketExtractor" ticket_app.py
"""

import os
import re
import sys
import json
import threading
import subprocess

try:
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    sys.exit("Run:  pip install pdfplumber openpyxl")

import tkinter as tk
from tkinter import filedialog, ttk, messagebox


# ============================ EXTRACTION CORE ===============================
AIRPORT_RE = r"[A-Z]{3}"
MONEY_RE = r"(?:INR|Rs\.?|₹|USD|\$)\s?[\d,]+(?:\.\d{1,2})?"
AIRLINES = {
    "indigo": "IndiGo", "6e": "IndiGo", "air india": "Air India",
    "vistara": "Vistara", "spicejet": "SpiceJet", "akasa": "Akasa Air",
    "go first": "Go First", "goair": "Go First", "emirates": "Emirates",
    "etihad": "Etihad", "qatar": "Qatar Airways", "lufthansa": "Lufthansa",
}
COLUMNS = ["Source File", "PNR", "Airline", "Flight No", "From", "To",
           "Travel Date", "Passenger", "Total Cost"]


def pdf_to_text(path):
    chunks = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            chunks.append(page.extract_text() or "")
    return "\n".join(chunks)


def first(patterns, text, flags=re.I):
    for p in patterns:
        m = re.search(p, text, flags)
        if m:
            return m.group(1).strip()
    return ""


def parse_regex(text, filename):
    low = text.lower()
    pnr = first([r"PNR\s*[:#-]?\s*([A-Z0-9]{5,8})",
                 r"Booking\s*Ref(?:erence)?\s*[:#-]?\s*([A-Z0-9]{5,8})",
                 r"Airline\s*PNR\s*[:#-]?\s*([A-Z0-9]{5,8})"], text)
    airline = next((name for key, name in AIRLINES.items() if key in low), "")
    frm = to = ""
    m = re.search(r"\b([A-Z]{3})\b[^A-Za-z]{0,6}(?:to|→|->|–)[^A-Za-z]{0,6}\b([A-Z]{3})\b", text)
    if m:
        frm, to = m.group(1), m.group(2)
    cost = first([r"Total\s*(?:Amount|Fare|Paid|Payable)?\s*[:#-]?\s*(" + MONEY_RE + ")",
                  r"Grand\s*Total\s*[:#-]?\s*(" + MONEY_RE + ")",
                  r"Amount\s*Paid\s*[:#-]?\s*(" + MONEY_RE + ")",
                  "(" + MONEY_RE + ")"], text)
    flight_no = first([r"Flight\s*(?:No\.?|Number)?\s*[:#-]?\s*([A-Z0-9]{2}\s?-?\s?\d{2,4})"], text)
    date = first([r"\b(\d{1,2}\s+[A-Za-z]{3,9}\s+\d{2,4})\b",
                  r"\b(\d{1,2}[/-]\d{1,2}[/-]\d{2,4})\b"], text)
    passenger = first([r"Passenger\s*(?:Name)?\s*[:#-]?\s*([A-Z][a-zA-Z .]+)",
                       r"Name\s*[:#-]?\s*(?:Mr|Ms|Mrs)?\.?\s*([A-Z][a-zA-Z .]+)"], text)
    return {"Source File": filename, "PNR": pnr, "Airline": airline, "Flight No": flight_no,
            "From": frm, "To": to, "Travel Date": date, "Passenger": passenger, "Total Cost": cost}


def write_excel(rows, out_path):
    wb = Workbook(); ws = wb.active; ws.title = "Flight Tickets"
    hfill = PatternFill("solid", fgColor="1F4E78"); hfont = Font(bold=True, color="FFFFFF")
    for c, col in enumerate(COLUMNS, 1):
        cell = ws.cell(1, c, col); cell.fill = hfill; cell.font = hfont
        cell.alignment = Alignment(horizontal="center")
    for r, row in enumerate(rows, 2):
        for c, col in enumerate(COLUMNS, 1):
            ws.cell(r, c, row.get(col, ""))
    for c, col in enumerate(COLUMNS, 1):
        w = max(len(col), max((len(str(row.get(col, ""))) for row in rows), default=10))
        ws.column_dimensions[ws.cell(1, c).column_letter].width = min(w + 3, 40)
    ws.freeze_panes = "A2"; wb.save(out_path)


# ================================ GUI =======================================
class App:
    def __init__(self, root):
        self.root = root
        self.files = []
        root.title("Flight Ticket Extractor")
        root.geometry("560x560")
        root.configure(bg="#F4F6F9")

        tk.Label(root, text="✈  Flight Ticket Extractor", font=("Segoe UI", 17, "bold"),
                 bg="#F4F6F9", fg="#1F4E78").pack(pady=(18, 2))
        tk.Label(root, text="Add your ticket PDFs, then create one Excel with all of them.",
                 font=("Segoe UI", 10), bg="#F4F6F9", fg="#555").pack()

        # drop / add zone
        zone = tk.Frame(root, bg="white", highlightbackground="#C3CEDD",
                        highlightthickness=2, bd=0)
        zone.pack(fill="both", expand=False, padx=24, pady=14, ipady=6)
        tk.Button(zone, text="➕  Add ticket PDFs", font=("Segoe UI", 12, "bold"),
                  bg="#1F4E78", fg="white", relief="flat", padx=16, pady=8,
                  command=self.add_files).pack(pady=12)
        self.count_lbl = tk.Label(zone, text="No files added yet",
                                  font=("Segoe UI", 9), bg="white", fg="#888")
        self.count_lbl.pack(pady=(0, 8))

        # file list
        self.listbox = tk.Listbox(root, height=8, font=("Consolas", 9),
                                  bg="white", fg="#222", selectbackground="#D6E4F0")
        self.listbox.pack(fill="both", expand=True, padx=24)
        bar = tk.Frame(root, bg="#F4F6F9"); bar.pack(pady=4)
        tk.Button(bar, text="Remove selected", command=self.remove_sel,
                  relief="flat", bg="#E6EAF0").pack(side="left", padx=4)
        tk.Button(bar, text="Clear all", command=self.clear,
                  relief="flat", bg="#E6EAF0").pack(side="left", padx=4)

        self.progress = ttk.Progressbar(root, mode="determinate")
        self.progress.pack(fill="x", padx=24, pady=(8, 2))
        self.status = tk.Label(root, text="", font=("Segoe UI", 9),
                               bg="#F4F6F9", fg="#1F4E78")
        self.status.pack()

        tk.Button(root, text="📊  Create Excel", font=("Segoe UI", 13, "bold"),
                  bg="#2E7D32", fg="white", relief="flat", padx=20, pady=10,
                  command=self.start).pack(pady=12)

    def add_files(self):
        paths = filedialog.askopenfilenames(title="Select ticket PDFs",
                                            filetypes=[("PDF files", "*.pdf")])
        for p in paths:
            if p not in self.files:
                self.files.append(p)
                self.listbox.insert("end", os.path.basename(p))
        self.refresh()

    def remove_sel(self):
        for i in reversed(self.listbox.curselection()):
            self.listbox.delete(i); self.files.pop(i)
        self.refresh()

    def clear(self):
        self.files.clear(); self.listbox.delete(0, "end"); self.refresh()

    def refresh(self):
        n = len(self.files)
        self.count_lbl.config(text=f"{n} file(s) ready" if n else "No files added yet")

    def start(self):
        if not self.files:
            messagebox.showwarning("No files", "Add at least one ticket PDF first.")
            return
        out = filedialog.asksaveasfilename(defaultextension=".xlsx",
                                           initialfile="flight_tickets.xlsx",
                                           filetypes=[("Excel", "*.xlsx")])
        if not out:
            return
        threading.Thread(target=self.run, args=(out,), daemon=True).start()

    def run(self, out):
        rows = []
        total = len(self.files)
        self.progress.config(maximum=total, value=0)
        for i, path in enumerate(self.files, 1):
            name = os.path.basename(path)
            self.status.config(text=f"Reading {name} ...")
            try:
                rows.append(parse_regex(pdf_to_text(path), name))
            except Exception as e:
                rows.append({"Source File": name, "PNR": "ERROR: " + str(e)[:50]})
            self.progress.config(value=i); self.root.update_idletasks()
        try:
            write_excel(rows, out)
        except Exception as e:
            messagebox.showerror("Error saving", str(e)); return
        self.status.config(text=f"Done — {total} ticket(s) saved.")
        if messagebox.askyesno("Success!",
                               f"Excel created with {total} ticket(s).\n\nOpen it now?"):
            self.open_file(out)

    @staticmethod
    def open_file(path):
        if sys.platform.startswith("win"):
            os.startfile(path)
        elif sys.platform == "darwin":
            subprocess.run(["open", path])
        else:
            subprocess.run(["xdg-open", path])


if __name__ == "__main__":
    root = tk.Tk()
    App(root)
    root.mainloop()
